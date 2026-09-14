from __future__ import annotations

import shutil
from copy import deepcopy
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "docs" / "SRSTemplate_v1.0.docx"
OUT_DIR = ROOT / "project doc"
OUT_DOCX = OUT_DIR / "PhishGuard_SRS.docx"
DIAGRAM_DIR = ROOT / "docs" / "diagrams"

NAVY = "17375E"
BLUE = "2F75B5"
TEAL = "008C82"
GOLD = "B7791F"
INK = "1F2937"
MUTED = "5B6777"
PALE_BLUE = "EEF4FA"
PALE_TEAL = "EAF6F3"
PALE_GOLD = "FFF7E6"
PALE_GRAY = "F5F7FA"
WHITE = "FFFFFF"
GRID = "D7DEE8"


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_run_font(run, name="Aptos", size=10.5, color=INK, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge not in kwargs:
            continue
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key in ["val", "sz", "space", "color"]:
            if key in kwargs[edge]:
                element.set(qn(f"w:{key}"), str(kwargs[edge][key]))


def set_table_geometry(table, widths_dxa: list[int], indent_dxa=120):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    set_run_font(run, size=9, color=MUTED)


def clear_body(doc: Document):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def clear_container(container):
    for p in list(container.paragraphs):
        p._element.getparent().remove(p._element)
    for table in list(container.tables):
        table._element.getparent().remove(table._element)
    container.add_paragraph()


def configure_styles(doc: Document):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.widow_control = True

    for name, size, color, before, after in [
        ("Heading 1", 16, NAVY, 15, 7),
        ("Heading 2", 12.5, BLUE, 11, 5),
        ("Heading 3", 11.5, TEAL, 8, 4),
    ]:
        s = styles[name]
        s.font.name = "Aptos Display" if name == "Heading 1" else "Aptos"
        s._element.rPr.rFonts.set(qn("w:ascii"), s.font.name)
        s._element.rPr.rFonts.set(qn("w:hAnsi"), s.font.name)
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = rgb(color)
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True
        s.paragraph_format.widow_control = True

    for name in ["List Bullet", "List Number"]:
        if name in [style.name for style in styles]:
            s = styles[name]
        else:
            s = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            s.base_style = normal
        s.font.name = "Aptos"
        s._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        s._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        s.font.size = Pt(10.5)
        s.font.color.rgb = rgb(INK)
        s.paragraph_format.left_indent = Inches(0.27)
        s.paragraph_format.first_line_indent = Inches(-0.16)
        s.paragraph_format.space_after = Pt(3)
        s.paragraph_format.line_spacing = 1.1

    if "Caption" not in [s.name for s in styles]:
        caption = styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        caption = styles["Caption"]
    caption.font.name = "Aptos"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = rgb(MUTED)
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(10)
    caption.paragraph_format.keep_with_next = False


def ensure_numbering(doc: Document):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    next_abstract = max(abstract_ids or [0]) + 1
    next_num = max(num_ids or [0]) + 1

    def add_abstract(abstract_id, fmt, text):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        lvl.append(start)
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        lvl.append(num_fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        lvl.append(lvl_text)
        jc = OxmlElement("w:lvlJc")
        jc.set(qn("w:val"), "left")
        lvl.append(jc)
        ppr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        tabs.append(tab)
        ppr.append(tabs)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "540")
        ind.set(qn("w:hanging"), "270")
        ppr.append(ind)
        lvl.append(ppr)
        if fmt == "bullet":
            rpr = OxmlElement("w:rPr")
            rfonts = OxmlElement("w:rFonts")
            rfonts.set(qn("w:ascii"), "Symbol")
            rfonts.set(qn("w:hAnsi"), "Symbol")
            rpr.append(rfonts)
            lvl.append(rpr)
        abstract.append(lvl)
        numbering.append(abstract)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(next_num))
        abs_ref = OxmlElement("w:abstractNumId")
        abs_ref.set(qn("w:val"), str(abstract_id))
        num.append(abs_ref)
        numbering.append(num)
        return next_num

    bullet_id = add_abstract(next_abstract, "bullet", "o")
    decimal_id = next_num + 1
    # The helper above uses the next available numId for the first list.
    # Build the second list explicitly so the ids remain deterministic.
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(next_abstract + 1))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1.")
    lvl.append(lvl_text)
    jc = OxmlElement("w:lvlJc")
    jc.set(qn("w:val"), "left")
    lvl.append(jc)
    ppr = OxmlElement("w:pPr")
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    ppr.append(ind)
    lvl.append(ppr)
    abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(decimal_id))
    abs_ref = OxmlElement("w:abstractNumId")
    abs_ref.set(qn("w:val"), str(next_abstract + 1))
    num.append(abs_ref)
    numbering.append(num)
    return bullet_id, decimal_id, next_abstract + 1


def new_num_id(doc: Document, abstract_id: int) -> int:
    numbering = doc.part.numbering_part.element
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    new_id = max(num_ids or [0]) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_id))
    abs_ref = OxmlElement("w:abstractNumId")
    abs_ref.set(qn("w:val"), str(abstract_id))
    num.append(abs_ref)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    numbering.append(num)
    return new_id


def apply_numbering(paragraph, num_id: int):
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = ppr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        ppr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_el)


def set_section_geometry(section):
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)


def configure_header_footer(section):
    clear_container(section.header)
    clear_container(section.footer)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    r = header.add_run("PHISHGUARD  |  SOFTWARE REQUIREMENTS SPECIFICATION")
    set_run_font(r, size=8.5, color=MUTED, bold=True)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    footer.paragraph_format.space_after = Pt(0)
    r = footer.add_run("Supervisor: Smita Dixit Kansal  |  Page ")
    set_run_font(r, size=8.5, color=MUTED)
    add_page_field(footer)


def add_para(doc, text="", style="Normal", align=None, before=None, after=None, keep=False):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = keep
    if text:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.left_indent = Inches(0.27)
        p.paragraph_format.first_line_indent = Inches(-0.16)
        p.paragraph_format.space_after = Pt(3)
        marker_run = p.add_run()
        symbol = OxmlElement("w:sym")
        symbol.set(qn("w:font"), "Symbol")
        symbol.set(qn("w:char"), "F0B7")
        marker_run._r.append(symbol)
        marker_run._r.append(OxmlElement("w:tab"))
        set_run_font(marker_run, name="Symbol", size=10, color=INK)
        r = p.add_run(item)
        set_run_font(r)


def add_numbered(doc, items):
    list_num_id = new_num_id(doc, doc._phishguard_decimal_abstract_id)
    for item in items:
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.left_indent = Inches(0.27)
        p.paragraph_format.first_line_indent = Inches(-0.16)
        p.paragraph_format.space_after = Pt(3)
        apply_numbering(p, list_num_id)
        r = p.add_run(item)
        set_run_font(r)


def add_table(doc, headers, rows, widths_dxa, header_fill=NAVY, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths_dxa)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, header_fill)
        set_cell_border(cell, top={"val": "single", "sz": 6, "color": GRID}, bottom={"val": "single", "sz": 6, "color": GRID}, left={"val": "single", "sz": 6, "color": GRID}, right={"val": "single", "sz": 6, "color": GRID})
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(header))
        set_run_font(r, size=font_size, color=WHITE, bold=True)
    mark_header_row(table.rows[0])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            if row_index % 2 == 1:
                set_cell_shading(cell, PALE_GRAY)
            set_cell_border(cell, top={"val": "single", "sz": 4, "color": GRID}, bottom={"val": "single", "sz": 4, "color": GRID}, left={"val": "single", "sz": 4, "color": GRID}, right={"val": "single", "sz": 4, "color": GRID})
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(str(value))
            set_run_font(r, size=font_size, color=INK)
    set_table_geometry(table, widths_dxa)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_callout(doc, label, text, fill=PALE_TEAL, accent=TEAL):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, top={"val": "single", "sz": 6, "color": accent}, bottom={"val": "single", "sz": 6, "color": accent}, left={"val": "single", "sz": 18, "color": accent}, right={"val": "single", "sz": 6, "color": accent})
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(label + ": ")
    set_run_font(r, size=10, color=accent, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=10, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_figure(doc, path: Path, caption: str, width_inches: float, alt_text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    shape = p.add_run().add_picture(str(path), width=Inches(width_inches))
    shape._inline.docPr.set("descr", alt_text)
    cap = doc.add_paragraph(style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    set_run_font(r, size=9, color=MUTED, italic=True)


def add_title_block(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(34)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("SOFTWARE REQUIREMENTS SPECIFICATION")
    set_run_font(r, size=10, color=TEAL, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("PhishGuard")
    set_run_font(r, name="Aptos Display", size=31, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(24)
    r = p.add_run("Intelligent Phishing Website Detection using Machine Learning,\nExplainable AI (SHAP), and FastAPI")
    set_run_font(r, size=14, color=MUTED)

    table = doc.add_table(rows=4, cols=2)
    set_table_geometry(table, [2200, 7160])
    meta = [
        ("Submitted to", "Smita Dixit Kansal\nAssistant Professor\nDepartment of Information Technology"),
        ("Submitted by", "Yash Gangal\nStudent"),
        ("Institution", "Institute of Technology and Science, Mohan Nagar, Ghaziabad"),
        ("Document", "Version 1.0  |  19 August 2026"),
    ]
    for i, (label, value) in enumerate(meta):
        label_cell, value_cell = table.rows[i].cells
        set_cell_shading(label_cell, PALE_BLUE)
        for cell in (label_cell, value_cell):
            set_cell_border(cell, top={"val": "single", "sz": 6, "color": GRID}, bottom={"val": "single", "sz": 6, "color": GRID}, left={"val": "single", "sz": 6, "color": GRID}, right={"val": "single", "sz": 6, "color": GRID})
        lp = label_cell.paragraphs[0]
        lp.paragraph_format.space_after = Pt(0)
        lr = lp.add_run(label)
        set_run_font(lr, size=10, color=NAVY, bold=True)
        vp = value_cell.paragraphs[0]
        vp.paragraph_format.space_after = Pt(0)
        vp.paragraph_format.line_spacing = 1.05
        lines = value.split("\n")
        for j, line in enumerate(lines):
            rr = vp.add_run(line)
            set_run_font(rr, size=10, color=INK, bold=(i == 0 and j == 0) or (i == 1 and j == 0))
            if j < len(lines) - 1:
                rr.add_break()
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    add_callout(doc, "Document purpose", "Defines the requirements, operating scenarios, interfaces, data contracts, quality targets, and acceptance criteria for the PhishGuard mini-project.", fill=PALE_GOLD, accent=GOLD)


def add_contents(doc):
    doc.add_heading("Contents", level=1)
    add_para(doc, "This contents list follows the section structure of the retained SRS template. Detailed page numbering is intentionally left to Word's field update when the document is opened.", after=8)
    contents = [
        "1. Purpose",
        "2. Scope",
        "3. Intended Audiences",
        "4. Definitions and Abbreviations",
        "5. References",
        "6. Assumptions and Dependencies",
        "7. Problem Statement",
        "8. Overall Description",
        "9. Use Cases",
        "10. Functional Requirements",
        "11. Data Requirements",
        "12. Non-Functional Requirements",
        "13. Activity Diagram",
        "14. Flow Chart",
        "15. Acceptance Criteria",
        "16. Risks, Constraints, and Future Work",
        "Appendix A - API Summary",
        "Appendix B - Feature Dictionary",
    ]
    for item in contents:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_run_font(r, size=10.5, color=NAVY, bold=item[0].isdigit() and "." in item[:3])


def build():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document(str(TEMPLATE))
    clear_body(doc)
    for section in doc.sections:
        set_section_geometry(section)
        configure_header_footer(section)
    configure_styles(doc)
    doc._phishguard_bullet_num_id, doc._phishguard_decimal_num_id, doc._phishguard_decimal_abstract_id = ensure_numbering(doc)
    doc.core_properties.title = "PhishGuard Software Requirements Specification"
    doc.core_properties.subject = "Requirements, scenarios, diagrams, data, and acceptance criteria for PhishGuard"
    doc.core_properties.author = "Yash Gangal"
    doc.core_properties.keywords = "PhishGuard, SRS, phishing detection, machine learning, FastAPI, SHAP"
    doc.core_properties.comments = "Prepared for Smita Dixit Kansal"

    add_title_block(doc)
    doc.add_page_break()
    add_contents(doc)
    doc.add_page_break()

    doc.add_heading("1. Purpose", level=1)
    add_para(doc, "This Software Requirements Specification (SRS) defines the functional, non-functional, data, interface, and operational requirements for PhishGuard. PhishGuard is a web-based phishing website detection system that accepts an absolute URL, extracts URL and optional webpage signals, produces an explainable machine-learning assessment, and stores an auditable scan record.")
    add_para(doc, "The document is intended for the supervisor, student developer, testers, reviewers, and future maintainers. It establishes the baseline for implementation, testing, demonstration, and academic evaluation.")

    doc.add_heading("2. Scope", level=1)
    doc.add_heading("2.1 In scope", level=2)
    add_bullets(doc, [
        "Validate and submit absolute HTTP(S) URLs through the user interface and FastAPI API.",
        "Extract 15 URL-structure features and 7 optional HTML/DOM features.",
        "Perform bounded best-effort webpage retrieval with URL-only degradation when retrieval fails.",
        "Run the active trained model or the configured deterministic heuristic fallback.",
        "Return phishing/legitimate verdict, confidence, 0-100 risk score, risk tier, and top contributing signals.",
        "Persist scan history and active-model metadata in SQLite.",
        "Support history review, filtering, pagination, deletion, analytics, model information, API documentation, and printable reports through the planned frontend.",
    ])
    doc.add_heading("2.2 Out of scope", level=2)
    add_bullets(doc, [
        "Blocking or quarantining webpages directly in a browser, proxy, or network gateway.",
        "Guaranteed detection of every zero-day or previously unseen phishing campaign.",
        "Paid reputation, blacklist, or third-party threat-intelligence API integration.",
        "Authentication, role-based access control, and multi-tenant isolation in the mini-project version.",
        "Automatic takedown, remediation, or notification of website owners.",
    ])

    doc.add_heading("3. Intended Audiences", level=1)
    add_table(doc, ["Audience", "Use of this document"], [
        ("Supervisor / reviewer", "Evaluate completeness, feasibility, traceability, and academic project alignment."),
        ("Student developer", "Use as the implementation and integration baseline."),
        ("Backend / frontend developers", "Implement API contracts, UI behavior, service boundaries, and data handling."),
        ("Testers", "Derive unit, integration, negative-path, performance, and acceptance tests."),
        ("Future maintainers", "Understand assumptions, risks, model behavior, and deferred work."),
    ], [2200, 7160])

    doc.add_heading("4. Definitions and Abbreviations", level=1)
    add_table(doc, ["Term", "Definition"], [
        ("SRS", "Software Requirements Specification."),
        ("URL", "Uniform Resource Locator submitted for analysis."),
        ("HTML / DOM", "Webpage markup and its parsed Document Object Model."),
        ("SHAP", "SHapley Additive exPlanations; a method used to rank feature contributions."),
        ("XAI", "Explainable Artificial Intelligence."),
        ("ML", "Machine Learning."),
        ("Risk score", "Integer from 0 to 100 derived from phishing probability."),
        ("URL-only analysis", "Analysis using the 15 URL features when HTML retrieval is unavailable."),
        ("Fallback model", "Deterministic heuristic classifier used when a trained artifact is unavailable and fallback is enabled."),
    ], [2200, 7160])

    doc.add_heading("5. References", level=1)
    add_numbered(doc, [
        "Retained SRS template: docs/SRSTemplate_v1.0.docx.",
        "PhishGuard project overview and architecture guide: OVERVIEW.md.",
        "System architecture: files/01_system_architecture.md.",
        "Database schema: files/02_database_schema.md.",
        "API contract: files/03_api_contract.md.",
        "Feature engineering specification: files/04_feature_engineering_spec.md.",
        "Model comparison report: phishing_detector/backend/ml/comparison_report.json.",
        "Backend implementation and tests: phishing_detector/backend/app and phishing_detector/backend/tests.",
    ])

    doc.add_heading("6. Assumptions and Dependencies", level=1)
    add_table(doc, ["Category", "Assumption / dependency"], [
        ("Runtime", "Python 3.11 is the supported backend runtime for the pinned scientific-ML dependencies."),
        ("Frontend", "The planned frontend uses React 18, Vite, Tailwind CSS, Chart.js, and Lucide icons."),
        ("Backend", "FastAPI, Uvicorn, Pydantic, SQLAlchemy, Requests, BeautifulSoup4, tldextract, scikit-learn, NumPy, and SHAP are available as configured."),
        ("Model artifact", "The trained artifact is optional for demonstration because the configured heuristic fallback can serve offline scans."),
        ("Network", "HTML enrichment depends on outbound access to the target URL and can fail or time out."),
        ("Persistence", "SQLite is appropriate for mini-project scale; production-scale concurrency may require a server database."),
        ("Security", "Fetching arbitrary URLs is treated as a security-sensitive operation and requires production SSRF controls beyond this baseline."),
    ], [2200, 7160])

    doc.add_heading("7. Problem Statement", level=1)
    add_para(doc, "Phishing websites imitate trusted services to collect credentials, payment information, or other sensitive data. Static blacklists can lag behind short-lived or newly generated domains. Paid reputation services can reduce implementation effort, but they introduce cost, network dependence, rate limits, and limited transparency into why a URL was flagged.")
    add_para(doc, "PhishGuard addresses this gap with a self-contained, explainable workflow that evaluates URL structure and optional webpage characteristics, returns a phishing or legitimate assessment, shows risk and evidence, and stores an auditable scan record. The workflow remains useful when the target site is offline, blocked, or too slow to fetch.")
    add_callout(doc, "Project goal", "Provide a fast, explainable, and auditable phishing-risk assessment without requiring paid threat-intelligence APIs.", fill=PALE_TEAL, accent=TEAL)
    doc.add_heading("7.1 Objectives", level=2)
    add_bullets(doc, [
        "Return a typed prediction through POST /predict after validating the submitted URL.",
        "Use a 22-feature hybrid vector: 15 always-available URL signals and 7 optional HTML signals.",
        "Use the benchmark-selected Random Forest model when its artifact is available, with a clearly identified heuristic fallback for offline demonstration.",
        "Provide confidence, a risk score from 0 to 100, low/medium/high risk tiers, and human-readable evidence.",
        "Persist the complete scan result with model version and UTC timestamp.",
        "Support analyst review through history, analytics, model metadata, and printable reports.",
    ])
    doc.add_heading("7.2 Scenario Document", level=2)
    add_para(doc, "Primary operating scenario - a user submits a URL for analysis:")
    add_numbered(doc, [
        "The user opens the Overview or Scan URL workspace and enters an absolute HTTP(S) URL.",
        "The frontend performs basic format validation and shows an inline error if the value is missing or malformed.",
        "The frontend sends the URL to POST /predict. FastAPI validates the request using the typed schema.",
        "The scraper makes a bounded best-effort request to the target page and counts redirects if the page responds.",
        "The feature service extracts 15 URL features and, when HTML is available, 7 HTML/DOM features. Otherwise it marks html_features_available=false and uses safe defaults for the HTML features.",
        "The model service loads the active model or permitted fallback, calculates phishing probability, prediction, confidence, risk score, and risk tier.",
        "The explainability service ranks the top five contributing signals using SHAP when supported or the deterministic fallback explanation otherwise.",
        "The database persists the complete result and model version in SQLite.",
        "The frontend displays verdict, confidence, risk, evidence, the full feature matrix, and an optional degradation note.",
        "The user can open the report, review the record later in history, or inspect aggregate analytics.",
    ])
    doc.add_heading("7.3 Exception Scenarios", level=2)
    add_table(doc, ["Condition", "Expected behavior"], [
        ("Malformed URL", "Return HTTP 422 with a human-readable validation detail; do not create a history row."),
        ("Target timeout, DNS, TLS, or HTTP failure", "Complete URL-only analysis, set html_features_available=false, and show the degradation note."),
        ("Model artifact missing", "Use the fallback when enabled and identify it in model metadata; otherwise return HTTP 503."),
        ("Unexpected extraction or persistence failure", "Rollback the transaction, log the failure, and return HTTP 500 with a generic detail."),
        ("Unknown history id", "Return HTTP 404 and leave existing records unchanged."),
        ("Concurrent scans", "Create independent records without partial or cross-linked results."),
    ], [2800, 6560])

    doc.add_heading("8. Overall Description", level=1)
    doc.add_heading("8.1 Product Perspective", level=2)
    add_para(doc, "PhishGuard is organized as a presentation layer, FastAPI API layer, feature services, ML/XAI services, and SQLite persistence. The planned React frontend communicates with the backend through JSON REST endpoints. The target website is an untrusted external input source used only for bounded enrichment.")
    add_table(doc, ["Layer", "Responsibility", "Representative technology"], [
        ("Presentation", "URL input, verdicts, history, analytics, model/API information, and report view.", "React 18, Vite, Tailwind CSS, Chart.js"),
        ("API", "Typed request/response contract, routing, validation, CORS, health, and orchestration.", "FastAPI, Pydantic, Uvicorn"),
        ("Feature services", "URL parsing, HTML retrieval, DOM parsing, and 22-feature aggregation.", "urllib, tldextract, Requests, BeautifulSoup4"),
        ("ML / XAI", "Model loading, probability inference, risk mapping, and top-five explanation.", "scikit-learn, Random Forest, SHAP, NumPy"),
        ("Persistence", "Scan history and active-model metadata.", "SQLite, SQLAlchemy, JSON"),
    ], [1900, 4200, 3260], font_size=8.8)
    doc.add_heading("8.2 Product Functions", level=2)
    add_bullets(doc, [
        "Quick scan from the overview and dedicated Scan URL workspace.",
        "Verdict display with confidence, risk score, risk tier, and top contributing signals.",
        "Full 22-feature matrix with HTML availability indicator.",
        "History search, filtering, pagination, detail view, deletion, and CSV export behavior.",
        "Analytics cards and model comparison visualization.",
        "Active model metadata, endpoint documentation, health status, and Swagger UI access.",
        "Print/PDF-ready detailed scan report.",
    ])
    doc.add_heading("8.3 User Characteristics", level=2)
    add_para(doc, "The primary user is an analyst, student, reviewer, or developer who can enter a URL and interpret a basic risk verdict. The interface must not require knowledge of feature engineering or model internals to understand the primary outcome, while still exposing enough evidence for technical review.")
    doc.add_heading("8.4 Operating Environment", level=2)
    add_bullets(doc, [
        "Development API: http://localhost:8000.",
        "FastAPI interactive documentation: http://localhost:8000/docs.",
        "Frontend development server: http://localhost:5173.",
        "SQLite database retained on local disk.",
        "Modern desktop browser for the planned responsive React interface.",
    ])
    doc.add_heading("8.5 Constraints", level=2)
    add_bullets(doc, [
        "The model output is a risk assessment, not a guarantee or a replacement for analyst judgment.",
        "HTML features are unavailable for unreachable targets and must be disclosed to the user.",
        "SQLite is not intended for high-volume production concurrency.",
        "The mini-project baseline has no authentication or access-control boundary.",
        "External page fetching requires production controls against SSRF, excessive response sizes, and unsafe network access.",
    ])

    doc.add_heading("9. Use Cases", level=1)
    doc.add_heading("9.1 Use Case Diagram", level=2)
    add_para(doc, "The diagram identifies the analyst/user as the primary actor and the target website as an optional enrichment source.")
    add_figure(doc, DIAGRAM_DIR / "PhishGuard_Use_Case_Diagram.png", "Figure 9-1. PhishGuard use case diagram.", 6.35, "Use case diagram showing Analyst/User and Target Website actors connected to PhishGuard capabilities.")
    doc.add_heading("9.2 Use Case Summary", level=2)
    add_table(doc, ["ID / use case", "Primary actor", "Precondition", "Success outcome"], [
        ("UC-01 Submit URL", "Analyst / User", "Scan page is available.", "A valid URL is sent to POST /predict."),
        ("UC-02 Review verdict", "Analyst / User", "A scan response exists.", "Prediction, confidence, risk, and evidence are visible."),
        ("UC-03 Review history", "Analyst / User", "Database is reachable.", "Paginated records are displayed and can be filtered."),
        ("UC-04 Manage history", "Analyst / User", "A history record is selected.", "Record is viewed, exported, or deleted after confirmation."),
        ("UC-05 Review analytics", "Analyst / User", "History/model data is available.", "Aggregate scan and model metrics are visible."),
        ("UC-06 Generate report", "Analyst / User", "A detailed scan is selected.", "A print/PDF-ready report is displayed."),
        ("UC-07 Enrich scan", "Target Website", "URL is syntactically valid.", "HTML features and redirect count are added when reachable."),
    ], [1800, 1800, 2500, 3260], font_size=8.8)
    doc.add_heading("9.3 Primary Use Case Narrative - UC-01 Submit URL", level=2)
    add_table(doc, ["Field", "Description"], [
        ("Primary actor", "Analyst / User"),
        ("Trigger", "User selects Scan after entering a URL."),
        ("Preconditions", "Frontend and API are available; the URL field is visible."),
        ("Main flow", "Validate input, call POST /predict, fetch page if enabled, extract features, infer, explain, persist, and render the result."),
        ("Alternate flow", "If the target cannot be fetched, continue with URL-only features and disclose the limitation."),
        ("Failure flow", "Return 422 for invalid input, 503 for unavailable model without fallback, or 500 for an unexpected server failure."),
        ("Postcondition", "A successful scan is stored and the user can review the verdict and evidence."),
    ], [2200, 7160])

    doc.add_heading("10. Functional Requirements", level=1)
    doc.add_heading("10.1 API Input and Validation", level=2)
    add_table(doc, ["ID", "Requirement", "Priority", "Verification"], [
        ("REQ-001", "The system shall expose GET /health and report service status, model readiness, and database connectivity.", "Must", "API test"),
        ("REQ-002", "The system shall accept a JSON request containing one required url field.", "Must", "Schema test"),
        ("REQ-003", "The system shall reject missing, malformed, or non-absolute URL values with HTTP 422.", "Must", "Negative API test"),
        ("REQ-004", "The system shall normalize the submitted URL to a string and persist the hostname/domain in lowercase.", "Must", "Integration test"),
        ("REQ-005", "The system shall return human-readable error details without exposing stack traces to the client.", "Must", "Error-path test"),
    ], [900, 5600, 900, 1960], font_size=8.8)
    doc.add_heading("10.2 Feature Engineering and Inference", level=2)
    add_table(doc, ["ID", "Requirement", "Priority", "Verification"], [
        ("REQ-006", "The system shall attempt a bounded webpage fetch using a configurable timeout and a PhishGuard user agent.", "Must", "Integration test"),
        ("REQ-007", "The system shall extract all 15 URL features for every syntactically valid URL without requiring a network call.", "Must", "Unit test"),
        ("REQ-008", "The system shall extract seven HTML/DOM features when page HTML is available and set html_features_available=true.", "Must", "Unit/integration test"),
        ("REQ-009", "When the target page cannot be fetched, the system shall use safe defaults for the seven HTML features, set html_features_available=false, and continue URL-only prediction.", "Must", "Failure-path test"),
        ("REQ-010", "The system shall load the active trained model when available or use the configured deterministic fallback when permitted.", "Must", "Configuration test"),
        ("REQ-011", "The system shall classify each scan as exactly phishing or legitimate based on the active model probability threshold.", "Must", "Unit/API test"),
        ("REQ-012", "The system shall return confidence as a decimal in the range 0.0 to 1.0 for the selected prediction.", "Must", "Schema test"),
        ("REQ-013", "The system shall calculate an integer risk score from 0 to 100 using the phishing probability.", "Must", "Unit test"),
        ("REQ-014", "The system shall map risk scores 0-33 to low, 34-66 to medium, and 67-100 to high.", "Must", "Boundary test"),
        ("REQ-015", "The system shall return the five largest absolute feature contributions using SHAP where supported and a deterministic fallback explanation otherwise.", "Must", "Unit/API test"),
    ], [900, 5600, 900, 1960], font_size=8.6)
    doc.add_heading("10.3 Persistence and History", level=2)
    add_table(doc, ["ID", "Requirement", "Priority", "Verification"], [
        ("REQ-016", "The system shall persist scan id, URL, domain, prediction, confidence, risk score, risk level, features JSON, top-features JSON, model version, and UTC timestamp.", "Must", "Database inspection"),
        ("REQ-017", "The system shall create or reuse metadata for the model version used by a scan, including metrics when available.", "Must", "Database inspection"),
        ("REQ-018", "The system shall return paginated history records in reverse chronological order.", "Must", "API test"),
        ("REQ-019", "The history endpoint shall support optional filtering for phishing or legitimate records.", "Must", "API test"),
        ("REQ-020", "The system shall delete an existing history record only when its identifier exists and shall return HTTP 404 for an unknown identifier.", "Must", "API test"),
        ("REQ-021", "The frontend shall provide a confirmation step before deleting a history record.", "Should", "UI demonstration"),
    ], [900, 5600, 900, 1960], font_size=8.8)
    doc.add_heading("10.4 Frontend, Analytics, and Reporting", level=2)
    add_table(doc, ["ID", "Requirement", "Priority", "Verification"], [
        ("REQ-022", "The Overview page shall expose health, active-model information, a quick scan input, and the latest scan summary.", "Should", "UI demonstration"),
        ("REQ-023", "The Scan URL workspace shall show validation, active-analysis state, verdict, confidence, risk score, risk level, and degradation status.", "Must", "UI demonstration"),
        ("REQ-024", "The result view shall display the top five contributing signals with value, impact, direction, and plain-language meaning where available.", "Must", "UI demonstration"),
        ("REQ-025", "The result view shall expose the full 22-feature matrix and indicate whether HTML enrichment was available.", "Must", "UI demonstration"),
        ("REQ-026", "The History page shall provide search/filter controls, pagination, detail viewing, deletion, and CSV export behavior.", "Should", "UI demonstration"),
        ("REQ-027", "The Analytics page shall display total scans, phishing count/percentage, average confidence, model accuracy, and comparison artifacts when available.", "Should", "UI demonstration"),
        ("REQ-028", "The Model and API page shall display active-model metadata, endpoint descriptions, health status, and Swagger UI access.", "Should", "UI demonstration"),
        ("REQ-029", "The report page shall present a print/PDF-ready audit view containing URL, verdict, risk, evidence, feature matrix, model signature, and timestamp.", "Should", "UI demonstration"),
    ], [900, 5600, 900, 1960], font_size=8.6)
    doc.add_heading("10.5 Reliability and API Documentation", level=2)
    add_table(doc, ["ID", "Requirement", "Priority", "Verification"], [
        ("REQ-030", "The backend shall log bounded scrape failures with diagnostic context without returning sensitive internal exception details to the client.", "Must", "Log inspection"),
        ("REQ-031", "The backend shall roll back the database transaction when prediction persistence fails.", "Must", "Failure-path test"),
        ("REQ-032", "The API shall expose OpenAPI/Swagger documentation for supported endpoints.", "Must", "Endpoint inspection"),
    ], [900, 5600, 900, 1960], font_size=8.8)

    doc.add_heading("11. Data Requirements", level=1)
    doc.add_heading("11.1 ScanHistory", level=2)
    add_table(doc, ["Field", "Type", "Required", "Description"], [
        ("id / scan_id", "Integer", "System", "Auto-incrementing identifier for the scan."),
        ("url", "VARCHAR(2048)", "Yes", "Validated URL submitted by the user."),
        ("domain", "VARCHAR(255)", "System", "Lowercase hostname extracted from the URL."),
        ("prediction", "VARCHAR(20)", "System", "phishing or legitimate."),
        ("confidence", "FLOAT", "System", "Confidence for the selected prediction, 0.0-1.0."),
        ("risk_score", "INTEGER", "System", "Phishing probability mapped to 0-100."),
        ("risk_level", "VARCHAR(10)", "System", "low, medium, or high."),
        ("features_json", "TEXT", "System", "22 named features plus html_features_available."),
        ("top_features_json", "TEXT", "System", "Top five contributing signals and directions."),
        ("model_version", "VARCHAR(50)", "System", "Model artifact or fallback version used for the scan."),
        ("scanned_at", "DATETIME", "System", "UTC timestamp for the completed scan."),
    ], [1900, 1600, 1000, 4860], font_size=8.6)
    doc.add_heading("11.2 ModelMetadata", level=2)
    add_table(doc, ["Field group", "Description"], [
        ("Identity", "id, model_name, algorithm, version, and is_active."),
        ("Evaluation", "accuracy, precision, recall, f1_score, and roc_auc when a benchmark report is available."),
        ("Training context", "dataset_size, trained_at, file_path, and optional notes."),
    ], [2200, 7160])
    doc.add_heading("11.3 Data Integrity Rules", level=2)
    add_bullets(doc, [
        "Prediction values are restricted to phishing or legitimate.",
        "Risk score is an integer between 0 and 100, and risk level follows the documented boundaries.",
        "Each successful scan references the model version that generated the result.",
        "History is ordered by scanned_at descending for the default API response.",
        "Features and top-feature explanations are stored as JSON so the complete evidence can be reconstructed for a report.",
    ])

    doc.add_heading("12. Non-Functional Requirements", level=1)
    add_table(doc, ["ID", "Category", "Requirement", "Verification"], [
        ("REQ-033", "Performance", "For URL-only analysis under normal local conditions, the API should return a prediction within 500 ms after request validation, excluding external page-fetch time.", "Performance test"),
        ("REQ-034", "Reliability", "The webpage fetch shall use a bounded timeout and shall not prevent a response when the target is offline or unreachable.", "Failure-path test"),
        ("REQ-035", "Compatibility", "The system shall use strict typed request/response schemas and shall not silently accept malformed URL input.", "Contract test"),
        ("REQ-036", "Architecture", "The baseline workflow shall avoid paid third-party threat-intelligence APIs.", "Architecture inspection"),
        ("REQ-037", "Usability", "The frontend shall present risk, confidence, evidence, and HTML availability with clear visual distinction and readable labels.", "UI inspection"),
        ("REQ-038", "Maintainability", "Routing, feature extraction, scraping, inference, explainability, persistence, and schemas shall remain separable modules.", "Code review"),
        ("REQ-039", "Durability", "Scan history shall persist across backend restarts when the SQLite database is retained.", "Restart test"),
        ("REQ-040", "Testability", "The application shall be testable with feature unit tests and API tests for health, prediction, history, validation, and deletion.", "Test run"),
        ("REQ-041", "Transparency", "The system shall disclose whether the result used full HTML enrichment or URL-only fallback.", "UI/API inspection"),
        ("REQ-042", "Security", "Production deployments shall apply SSRF protections, response-size limits, outbound network restrictions, and safe logging around target-page fetches.", "Security review"),
    ], [900, 1600, 4900, 1960], font_size=8.3)

    doc.add_heading("13. Activity Diagram", level=1)
    add_para(doc, "The activity flow below describes the request lifecycle from URL entry through validation, optional enrichment, feature extraction, inference, explanation, persistence, and UI rendering.")
    add_figure(doc, DIAGRAM_DIR / "PhishGuard_Activity_Diagram.png", "Figure 13-1. PhishGuard activity diagram for a URL scan request.", 4.85, "Activity diagram showing URL entry, validation, optional HTML fetch, feature extraction, model inference, explanation, persistence, and response rendering.")

    doc.add_heading("14. Flow Chart", level=1)
    add_para(doc, "The flow chart presents the end-to-end data movement across the React UI, FastAPI endpoint, feature service, ML/XAI service, and SQLite-backed response path, including key HTTP outcomes.")
    add_figure(doc, DIAGRAM_DIR / "PhishGuard_Flow_Chart.png", "Figure 14-1. PhishGuard scan flow chart.", 6.25, "Flow chart showing user URL submission, frontend validation, FastAPI prediction, feature engineering, ML and SHAP, SQLite response, optional HTML branch, and error outcomes.")

    doc.add_heading("15. Acceptance Criteria", level=1)
    add_table(doc, ["ID", "Acceptance criterion", "Evidence"], [
        ("AC-01", "A valid URL returns HTTP 200 with scan id, prediction, confidence, risk score, risk level, features, top features, model version, and timestamp.", "API test response"),
        ("AC-02", "The response contains exactly 22 feature values plus html_features_available.", "Schema/unit test"),
        ("AC-03", "An unreachable target still returns a URL-only prediction and clearly marks HTML enrichment unavailable.", "Integration/failure test"),
        ("AC-04", "An invalid URL returns HTTP 422 and creates no scan-history record.", "Negative API test"),
        ("AC-05", "A prediction is stored and can be retrieved from paginated history.", "Database and API test"),
        ("AC-06", "An existing record can be deleted and a missing record returns HTTP 404.", "API test"),
        ("AC-07", "The dashboard can present verdict, risk, confidence, evidence, and the full feature matrix.", "UI demonstration"),
        ("AC-08", "The SRS package contains the problem statement, scenario document, use case diagram, activity diagram, and flow chart aligned to implementation behavior.", "Document review"),
    ], [900, 6260, 2160], font_size=8.6)

    doc.add_heading("16. Risks, Constraints, and Future Work", level=1)
    add_table(doc, ["Risk / constraint", "Consequence", "Mitigation / future work"], [
        ("False positive or false negative", "A legitimate site may be flagged or a phishing site may be missed.", "Show confidence and evidence, keep the user in the review loop, retrain with representative data, and monitor metrics."),
        ("HTML fetch / SSRF exposure", "Fetching arbitrary URLs can consume resources or reach unintended internal services.", "Apply timeouts, response-size limits, URL/IP policy checks, egress restrictions, and production SSRF controls."),
        ("Feature drift", "Attackers can change URL and page construction patterns.", "Track history, review explanations, refresh datasets, and compare candidate models periodically."),
        ("SQLite concurrency limits", "High concurrent writes can cause contention.", "Retain SQLite for mini-project scale; migrate to a production database for larger deployments."),
        ("Fallback model overconfidence", "Heuristic predictions may look authoritative without a trained artifact.", "Identify fallback mode in model metadata and reports; never present fallback values as trained-model metrics."),
        ("No authentication in baseline", "Anyone with network access can submit URLs and view/delete local history.", "Add authentication, authorization, audit logging, and deployment isolation in a future release."),
    ], [2500, 3000, 3860], font_size=8.5)

    doc.add_page_break()
    doc.add_heading("Appendix A - API Summary", level=1)
    add_table(doc, ["Method / path", "Purpose", "Important response"], [
        ("GET /health", "Liveness/readiness check.", "status, model_loaded, database_connected"),
        ("POST /predict", "Validate, analyze, explain, persist, and return one URL scan.", "scan_id, prediction, confidence, risk_score, risk_level, features, top_features, html_features_available"),
        ("GET /history", "Return paginated history with optional prediction filter.", "items and pagination"),
        ("DELETE /history/{id}", "Delete one existing scan history record.", "204 on success; 404 if missing"),
        ("GET /model-info", "Return active model metadata and evaluation metrics.", "model name/version, metrics, dataset size, trained_at"),
        ("GET /docs", "FastAPI-generated interactive OpenAPI/Swagger UI.", "Interactive endpoint documentation"),
    ], [2200, 3600, 3560], font_size=8.7)
    doc.add_heading("Appendix B - Feature Dictionary", level=1)
    feature_rows = [
        ("1", "url_length", "URL", "Length of submitted URL."),
        ("2", "domain_length", "URL", "Length of extracted registrable domain."),
        ("3", "num_dots", "URL", "Count of dots across the URL."),
        ("4", "num_hyphens", "URL", "Count of hyphens."),
        ("5", "num_digits", "URL", "Count of decimal digits."),
        ("6", "num_subdomains", "URL", "Number of subdomain labels."),
        ("7", "has_https", "URL", "Whether the scheme is HTTPS."),
        ("8", "has_ip_address", "URL", "Whether hostname is an IP literal."),
        ("9", "has_at_symbol", "URL", "Whether URL contains @ user-info syntax."),
        ("10", "has_double_slash_redirect", "URL", "Whether an extra // appears after the scheme."),
        ("11", "is_shortened_url", "URL", "Whether hostname matches a known shortener."),
        ("12", "num_suspicious_chars", "URL", "Count of @, %, _, =, and & obfuscation characters."),
        ("13", "url_entropy", "URL", "Shannon entropy of URL characters."),
        ("14", "has_suspicious_tld", "URL", "Whether suffix is in the configured abused-TLD set."),
        ("15", "path_length", "URL", "Length of parsed URL path."),
        ("16", "has_iframe", "HTML", "Whether HTML contains an iframe."),
        ("17", "redirect_count", "HTML", "Number of HTTP redirects followed."),
        ("18", "num_external_links", "HTML", "Count of anchor links to other hostnames."),
        ("19", "form_action_suspicious", "HTML", "Whether a form action is empty, external, or an IP literal."),
        ("20", "has_javascript_events", "HTML", "Whether inline event handlers are present."),
        ("21", "has_popup_window", "HTML", "Whether JavaScript calls window.open()."),
        ("22", "has_hidden_elements", "HTML", "Whether inline styles hide elements."),
        ("-", "html_features_available", "Metadata", "Whether the seven HTML features were computed from fetched page content."),
    ]
    add_table(doc, ["#", "Feature", "Group", "Definition"], feature_rows, [500, 2500, 1100, 5260], font_size=8.4)
    add_callout(doc, "Model baseline", "The repository comparison report records RandomForest as the selected model because it has the highest F1 and ROC-AUC in the current benchmark: accuracy 0.9195, precision 0.8992, recall 0.8873, F1 0.8932, ROC-AUC 0.9693, dataset size 1,225,480. These are benchmark values, not a guarantee for every future dataset or URL population.", fill=PALE_BLUE, accent=BLUE)

    # Ensure fields are refreshed when the document is opened in Word.
    settings = doc.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")
    doc.save(str(OUT_DOCX))
    return OUT_DOCX


if __name__ == "__main__":
    out = build()
    print(out)
